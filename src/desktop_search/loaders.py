"""File loaders. One function per file type plus a dispatcher.

Each loader returns ``list[Document]``. A multi-page/sheet/slide file
produces one ``Document`` per logical section; whitespace-only sections
are dropped. Single-document formats produce zero or one ``Document``.
Unknown extensions return an empty list rather than raising.
"""

from pathlib import Path
from typing import Callable, NamedTuple

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


class Document(NamedTuple):
    text: str
    source: Path
    section: str | None
    mtime: float


PLAIN_TEXT_EXTS: frozenset[str] = frozenset({
    ".md", ".markdown", ".txt", ".rst",
    ".html", ".htm", ".css", ".js", ".ts", ".jsx", ".tsx",
    ".py", ".rb", ".php", ".java", ".kt", ".swift", ".m", ".mm",
    ".c", ".cc", ".cpp", ".h", ".hpp", ".rs", ".go", ".cs",
    ".yaml", ".yml", ".json", ".toml", ".ini", ".cfg", ".sh", ".bash", ".zsh",
    ".sql", ".xml",
})


def _mtime(path: Path) -> float:
    return path.stat().st_mtime


def load_pdf(path: Path) -> list[Document]:
    mtime = _mtime(path)
    reader = PdfReader(str(path))
    docs: list[Document] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(text=text, source=path, section=f"p.{i}", mtime=mtime))
    return docs


def load_docx(path: Path) -> list[Document]:
    mtime = _mtime(path)
    doc = DocxDocument(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()
    if not text:
        return []
    return [Document(text=text, source=path, section=None, mtime=mtime)]


def load_pptx(path: Path) -> list[Document]:
    mtime = _mtime(path)
    pres = Presentation(str(path))
    docs: list[Document] = []
    for i, slide in enumerate(pres.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                if para.text.strip():
                    lines.append(para.text)
        text = "\n".join(lines).strip()
        if text:
            docs.append(Document(text=text, source=path, section=f"slide {i}", mtime=mtime))
    return docs


def load_xlsx(path: Path) -> list[Document]:
    mtime = _mtime(path)
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        docs: list[Document] = []
        for sheet in wb.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append("\t".join(cells))
            text = "\n".join(rows).strip()
            if text:
                docs.append(Document(text=text, source=path, section=sheet.title, mtime=mtime))
        return docs
    finally:
        wb.close()


def load_plain_text(path: Path) -> list[Document]:
    mtime = _mtime(path)
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    return [Document(text=text, source=path, section=None, mtime=mtime)]


_LOADERS: dict[str, Callable[[Path], list[Document]]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".pptx": load_pptx,
    ".xlsx": load_xlsx,
}


def load_file(path: Path) -> list[Document]:
    ext = path.suffix.lower()
    if ext in _LOADERS:
        return _LOADERS[ext](path)
    if ext in PLAIN_TEXT_EXTS:
        return load_plain_text(path)
    return []
