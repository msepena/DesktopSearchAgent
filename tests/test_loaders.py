from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF
from openpyxl import Workbook
from pptx import Presentation

from desktop_search.loaders import (
    load_docx,
    load_file,
    load_pdf,
    load_plain_text,
    load_pptx,
    load_xlsx,
)


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    p = tmp_path / "sample.pdf"
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    pdf.add_page()
    pdf.cell(text="Hello PDF page one")
    pdf.add_page()
    pdf.cell(text="Goodbye PDF page two")
    pdf.output(str(p))
    return p


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    p = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("Hello Word document")
    doc.add_paragraph("Second paragraph here")
    doc.save(str(p))
    return p


@pytest.fixture
def pptx_file(tmp_path: Path) -> Path:
    p = tmp_path / "sample.pptx"
    pres = Presentation()
    s1 = pres.slides.add_slide(pres.slide_layouts[5])
    s1.shapes.title.text = "Slide One Title"
    s2 = pres.slides.add_slide(pres.slide_layouts[5])
    s2.shapes.title.text = "Slide Two Title"
    pres.save(str(p))
    return p


@pytest.fixture
def xlsx_file(tmp_path: Path) -> Path:
    p = tmp_path / "sample.xlsx"
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Data"
    sheet["A1"] = "Name"
    sheet["B1"] = "Score"
    sheet["A2"] = "Alice"
    sheet["B2"] = 95
    notes = wb.create_sheet(title="Notes")
    notes["A1"] = "Random note"
    wb.save(str(p))
    return p


@pytest.fixture
def md_file(tmp_path: Path) -> Path:
    p = tmp_path / "notes.md"
    p.write_text("# Heading\n\nSome markdown body text.\n")
    return p


def test_load_pdf_one_doc_per_page(pdf_file: Path) -> None:
    docs = load_pdf(pdf_file)
    assert len(docs) == 2
    assert docs[0].section == "p.1"
    assert docs[1].section == "p.2"
    assert "Hello" in docs[0].text
    assert "Goodbye" in docs[1].text


def test_load_docx_returns_single_doc(docx_file: Path) -> None:
    docs = load_docx(docx_file)
    assert len(docs) == 1
    assert "Hello Word" in docs[0].text
    assert "Second paragraph" in docs[0].text
    assert docs[0].section is None


def test_load_pptx_one_doc_per_slide(pptx_file: Path) -> None:
    docs = load_pptx(pptx_file)
    assert len(docs) == 2
    assert docs[0].section == "slide 1"
    assert "Slide One Title" in docs[0].text
    assert docs[1].section == "slide 2"
    assert "Slide Two Title" in docs[1].text


def test_load_xlsx_one_doc_per_sheet(xlsx_file: Path) -> None:
    docs = load_xlsx(xlsx_file)
    assert len(docs) == 2
    sections = {d.section for d in docs}
    assert sections == {"Data", "Notes"}
    data_doc = next(d for d in docs if d.section == "Data")
    assert "Alice" in data_doc.text
    assert "95" in data_doc.text


def test_load_plain_text(md_file: Path) -> None:
    docs = load_plain_text(md_file)
    assert len(docs) == 1
    assert "markdown body" in docs[0].text


def test_dispatcher_routes_by_extension(
    md_file: Path, pdf_file: Path, docx_file: Path, pptx_file: Path, xlsx_file: Path
) -> None:
    assert len(load_file(md_file)) == 1
    assert len(load_file(pdf_file)) == 2
    assert len(load_file(docx_file)) == 1
    assert len(load_file(pptx_file)) == 2
    assert len(load_file(xlsx_file)) == 2


def test_unknown_extension_returns_empty(tmp_path: Path) -> None:
    p = tmp_path / "blob.bin"
    p.write_bytes(b"\x00\x01\x02")
    assert load_file(p) == []


def test_documents_carry_source_and_mtime(md_file: Path) -> None:
    docs = load_file(md_file)
    assert docs[0].source == md_file
    assert docs[0].mtime > 0


def test_empty_plain_text_returns_empty(tmp_path: Path) -> None:
    p = tmp_path / "blank.md"
    p.write_text("   \n\n  ")
    assert load_plain_text(p) == []
